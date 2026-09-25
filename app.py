from flask import Flask, request, jsonify, send_file, render_template
import arabic_reshaper
from bidi.algorithm import get_display
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from pypdf import PdfWriter, PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json, os, io, base64, tempfile
import zipfile, unicodedata, re
from pathlib import Path

app = Flask(__name__)
BASE = Path(__file__).parent

# Dossiers d'export PDF par niveau
EXPORTS_DIR = BASE / 'exports' / 'pdf'
for _n in [1, 2, 3]:
    (EXPORTS_DIR / f'niv{_n}').mkdir(parents=True, exist_ok=True)

def safe_filename(nom):
    """Convertit un nom en nom de fichier safe ASCII."""
    nom = unicodedata.normalize('NFKD', str(nom)).encode('ascii', 'ignore').decode()
    nom = re.sub(r'[^\w\s-]', '', nom).strip()
    nom = re.sub(r'[\s]+', '_', nom)
    return nom or 'eleve'

# ── Fonts ──────────────────────────────────────────────────
FONTS_DIR = BASE / 'static' / 'fonts'
AMIRI     = str(FONTS_DIR / 'Amiri-Regular.ttf')
AMIRI_B   = str(FONTS_DIR / 'Amiri-Bold.ttf')
DEJAVU    = str(FONTS_DIR / 'DejaVuSans.ttf')
DEJAVU_B  = str(FONTS_DIR / 'DejaVuSans-Bold.ttf')
pdfmetrics.registerFont(TTFont('Amiri',     AMIRI))
pdfmetrics.registerFont(TTFont('AmiriBold', AMIRI_B))
pdfmetrics.registerFont(TTFont('DejaVu',    DEJAVU))
pdfmetrics.registerFont(TTFont('DejaVuBold',DEJAVU_B))

LOGO_PATH  = str(BASE / 'static' / 'logo.png')
STAMP_PATH = str(BASE / 'static' / 'stamp.png')

# ── Logo embarqué dans la page d'accueil ───────────────────
# Le logo de l'école est inséré directement dans le HTML (data URI) plutôt que
# chargé par une seconde requête : c'est cette requête qui échouait au réveil de
# l'hébergeur ou sur une connexion lente, et le logo « disparaissait ». Inséré
# dans la page, il s'affiche dès que la page s'affiche.

def logo_en_data_uri(chemin, cote=192):
    """Renvoie le logo réduit en data URI, ou une chaîne vide s'il est illisible."""
    try:
        from PIL import Image
        with Image.open(chemin) as img:
            img = img.convert('RGBA')
            img.thumbnail((cote, cote), Image.LANCZOS)
            # Palette de 128 couleurs : l'emblème reste net et la page ne
            # s'alourdit que d'une dizaine de kilo-octets.
            img = img.convert('P', palette=Image.ADAPTIVE, colors=128)
            tampon = io.BytesIO()
            img.save(tampon, format='PNG', optimize=True)
            donnees = tampon.getvalue()
    except Exception:
        try:
            donnees = Path(chemin).read_bytes()
        except OSError:
            return ''
    return 'data:image/png;base64,' + base64.b64encode(donnees).decode('ascii')

LOGO_URI = logo_en_data_uri(LOGO_PATH)

# ── Palette par niveau ─────────────────────────────────────
PALETTES = {
    1: {'main': colors.HexColor('#2e7d32'), 'dark': colors.HexColor('#1b5e20'), 'th': colors.HexColor('#388e3c')},
    2: {'main': colors.HexColor('#6a1b9a'), 'dark': colors.HexColor('#4a148c'), 'th': colors.HexColor('#7b1fa2')},
    3: {'main': colors.HexColor('#1a2e5a'), 'dark': colors.HexColor('#0d1a38'), 'th': colors.HexColor('#2c4170')},
}
GOLD = colors.HexColor('#c8a84b')
WHITE = colors.white
BLK   = colors.black
LGRY  = colors.HexColor('#eef0f4')

W, H   = A4
MARGIN = 18
TW     = W - 2*MARGIN
W_OBS=0.36; W_NOTE=0.10; W_FR=0.30; W_AR=0.24
X0=MARGIN; X1=MARGIN+TW*W_OBS; X2=X1+TW*W_NOTE; X3=X2+TW*W_FR; X4=MARGIN+TW
COL_OBS=TW*W_OBS; COL_NOTE=TW*W_NOTE; COL_FR=TW*W_FR; COL_AR=TW*W_AR
PAD=4; RH_TH=22; RH_SEC=20; RH_AVG=29; RH_MG=32

# ── Matières par niveau ────────────────────────────────────
MATIERES = {
    1: {
        'arabe': [('الْقِرَاءَةُ وَالْفَهْمُ','Lecture & Compréhension'),('التَّعْبِيرُ الشَّفْوِيُّ','Expression Orale'),('الْإِمْلَاءُ','Dictée (Imlaa)'),('الْخَطُّ',"L'écriture")],
        'islami': [('التَّرْبِيَةُ الْإِسْلَامِيَّةُ','Éducation Islamique'),('الْقُرْآنُ الْكَرِيمُ','Noble Coran')],
    },
    2: {
        'arabe': [('الْقِرَاءَةُ وَالْفَهْمُ','Lecture & Compréhension'),('التَّعْبِيرُ الشَّفْوِيُّ','Expression Orale'),('التَّعْبِيرُ الْكِتَابِيُّ','Expression Écrite'),('الْإِمْلَاءُ','Dictée (Imlaa)'),('الْخَطُّ',"L'écriture")],
        'islami': [('التَّرْبِيَةُ الْإِسْلَامِيَّةُ','Éducation Islamique'),('الْقُرْآنُ الْكَرِيمُ','Noble Coran')],
    },
    3: {
        'arabe': [('الْقِرَاءَةُ وَالْفَهْمُ','Lecture & Compréhension'),('التَّعْبِيرُ الشَّفْوِيُّ','Expression Orale'),('التَّعْبِيرُ الْكِتَابِيُّ','Expression Écrite'),('الْإِمْلَاءُ','Dictée (Imlaa)'),('الْخَطُّ',"L'écriture")],
        'islami': [('التَّرْبِيَةُ الْإِسْلَامِيَّةُ','Éducation Islamique'),('التِّلَاوَةُ وَقَوَاعِدُ التَّجْوِيدِ','Tilawa & Règles du Tajwid')],
    },
}

def ar(t): return get_display(arabic_reshaper.reshape(str(t)))

def wrap_text(cv, txt, font, size, max_w):
    if not txt: return []
    cv.setFont(font, size)
    words = txt.split(); lines = []; line = ""
    for w in words:
        test = (line+" "+w).strip()
        if cv.stringWidth(test, font, size) <= max_w: line = test
        else:
            if line: lines.append(line)
            line = w
    if line: lines.append(line)
    return lines

def obs_rh(cv, obs, size=8):
    if not obs: return 20
    parts = obs.split('\n', 1)
    fr_lines = wrap_text(cv, parts[0], 'DejaVu', size, COL_OBS-2*PAD) if parts[0] else []
    ar_lines = wrap_text(cv, parts[1], 'Amiri', size+1, COL_OBS-2*PAD) if len(parts)>1 and parts[1] else []
    total = len(fr_lines) + len(ar_lines)
    return max(20, total*(size+2.5)+7)

def draw_obs(cv, obs, xl, yt, rh, size=8):
    if not obs: return
    parts = obs.split('\n', 1)
    fr_text = parts[0]
    ar_text = parts[1] if len(parts)>1 else ''
    fr_lines = wrap_text(cv, fr_text, 'DejaVu', size, COL_OBS-2*PAD) if fr_text else []
    ar_lines = wrap_text(cv, ar_text, 'Amiri', size+1, COL_OBS-2*PAD) if ar_text else []
    lh = size+2.5
    total = (len(fr_lines)+len(ar_lines))*lh
    y0 = yt-(rh-total)/2-size
    cv.setFillColor(BLK)
    for i, ln in enumerate(fr_lines):
        cv.setFont('DejaVu', size); cv.drawString(xl+PAD, y0-i*lh, ln)
    y_ar = y0 - len(fr_lines)*lh
    for i, ln in enumerate(ar_lines):
        cv.setFont('Amiri', size+1); cv.drawRightString(xl+COL_OBS-PAD, y_ar-i*lh, ar(ln))

def fit_sz(cv, txt, font, mx, mn, mw):
    sz = mx; cv.setFont(font, sz)
    while sz >= mn and cv.stringWidth(txt, font, sz) > mw: sz -= 0.5
    return sz

def draw_bulletin(cv, data, niveau):
    P = PALETTES[niveau]
    MAIN=P['main']; DARK=P['dark']; TH=P['th']
    mats = MATIERES[niveau]

    def cl(cy, rh):
        cv.setStrokeColor(MAIN); cv.setLineWidth(0.5)
        for x in [X1,X2,X3]: cv.line(x, cy-rh, x, cy)

    def th_row(cy):
        cv.setFillColor(TH); cv.setStrokeColor(MAIN); cv.setLineWidth(0.8)
        cv.rect(X0,cy-RH_TH,TW,RH_TH,fill=1); cl(cy,RH_TH); cv.setFillColor(WHITE)
        for txt,fnt,sz,cx in [(ar('مُلاحَظَات'),'Amiri',11,(X0+X1)/2),(ar('مُعَدَّل'),'Amiri',11,(X1+X2)/2),('Matière','DejaVuBold',10,(X2+X3)/2),(ar('المادة'),'Amiri',11,(X3+X4)/2)]:
            cv.setFont(fnt,sz); cv.drawCentredString(cx,cy-RH_TH+7,txt)
        return cy-RH_TH

    def sec_row(cy, arl, frl):
        cv.setFillColor(MAIN); cv.setStrokeColor(MAIN); cv.setLineWidth(0.5)
        cv.rect(X0,cy-RH_SEC,TW,RH_SEC,fill=1)
        cv.setFont('AmiriBold',11); cv.setFillColor(WHITE)
        cv.drawCentredString(W/2,cy-RH_SEC+5,ar(arl)+'   •   '+frl)
        return cy-RH_SEC

    def data_row(cy, arm, frm, note='', obs='', bg=WHITE):
        rh = obs_rh(cv, obs); mid = cy-rh/2
        cv.setFillColor(bg); cv.setStrokeColor(MAIN); cv.setLineWidth(0.4)
        cv.rect(X0,cy-rh,TW,rh,fill=1); cl(cy,rh)
        draw_obs(cv, obs, X0, cy, rh)
        if note:
            sz = fit_sz(cv,str(note),'DejaVuBold',13,7,COL_NOTE-2*PAD)
            cv.setFont('DejaVuBold',sz); cv.setFillColor(BLK)
            cv.drawCentredString((X1+X2)/2,mid-sz/2+1,str(note))
        frl2 = wrap_text(cv,frm,'DejaVu',10,COL_FR-2*PAD)
        lhf=13; totf=len(frl2)*lhf; yf=cy-(rh-totf)/2-10
        cv.setFont('DejaVu',10); cv.setFillColor(BLK)
        for i,ln in enumerate(frl2): cv.drawRightString(X3-PAD,yf-i*lhf,ln)
        cv.setFont('Amiri',12); cv.setFillColor(BLK)
        cv.drawCentredString((X3+X4)/2,mid-5,ar(arm))
        return cy-rh

    def avg_row(cy, arm, frm, val):
        cv.setFillColor(DARK); cv.setStrokeColor(MAIN); cv.setLineWidth(0.8)
        cv.rect(X0,cy-RH_AVG,TW,RH_AVG,fill=1); cl(cy,RH_AVG)
        cv.setFillColor(GOLD); cv.rect(X1,cy-RH_AVG,COL_NOTE,RH_AVG,fill=1,stroke=0)
        cv.setStrokeColor(MAIN); cv.setLineWidth(0.5); cv.rect(X1,cy-RH_AVG,COL_NOTE,RH_AVG,fill=0)
        s = f'{val:.2f} / 20' if isinstance(val,float) else '— / 20'
        sz = fit_sz(cv,s,'DejaVuBold',10,6,COL_NOTE-2*PAD)
        cv.setFont('DejaVuBold',sz); cv.setFillColor(WHITE)
        cv.drawCentredString((X1+X2)/2,cy-RH_AVG/2-sz/2+1,s)
        cv.setFillColor(WHITE); cv.setFont('DejaVuBold',10); cv.drawRightString(X3-PAD,cy-RH_AVG+9,frm)
        cv.setFont('AmiriBold',10); cv.drawCentredString((X3+X4)/2,cy-RH_AVG+10,ar(arm))
        return cy-RH_AVG

    def mg_row(cy, val):
        cv.setFillColor(DARK); cv.setStrokeColor(MAIN); cv.setLineWidth(1.2)
        cv.rect(X0,cy-RH_MG,TW,RH_MG,fill=1); cl(cy,RH_MG)
        cv.setFillColor(GOLD); cv.rect(X1,cy-RH_MG,COL_NOTE,RH_MG,fill=1,stroke=0)
        cv.setStrokeColor(MAIN); cv.setLineWidth(0.5); cv.rect(X1,cy-RH_MG,COL_NOTE,RH_MG,fill=0)
        s = f'{val:.2f} / 20' if isinstance(val,float) else '— / 20'
        sz = fit_sz(cv,s,'DejaVuBold',11,6,COL_NOTE-2*PAD)
        cv.setFont('DejaVuBold',sz); cv.setFillColor(WHITE)
        cv.drawCentredString((X1+X2)/2,cy-RH_MG/2-sz/2+1,s)
        cv.setFillColor(WHITE); cv.setFont('DejaVuBold',11); cv.drawRightString(X3-PAD,cy-RH_MG+10,'MOYENNE GÉNÉRALE')
        cv.setFont('AmiriBold',11); cv.drawCentredString((X3+X4)/2,cy-RH_MG+11,ar('الْمُعَدَّلُ الْعَامُّ'))
        return cy-RH_MG

    y = H-MARGIN
    logo = ImageReader(LOGO_PATH)
    cv.drawImage(logo,(W-115)/2,y-86,width=115,height=86,mask='auto'); y -= 88

    HDR=100; cv.setStrokeColor(MAIN); cv.setLineWidth(2); cv.rect(MARGIN,y-HDR,TW,HDR)
    cv.setFont('Amiri',13); cv.setFillColor(GOLD)
    cv.drawCentredString(W/2,y-16,ar('بِسْمِ اللهِ الرَّحْمَنِ الرَّحِيمِ'))
    cv.setFont('AmiriBold',16); cv.setFillColor(MAIN)
    cv.drawCentredString(W/2,y-34,ar('مَدْرَسَةُ مَسْجِدِ كَاسْتِيلْ سَارَزَان'))
    cv.setFont('DejaVuBold',15); cv.drawCentredString(W/2,y-52,'École Mosquée de Castelsarrasin')
    cv.setFont('AmiriBold',13)
    cv.drawCentredString(W/2,y-68,ar(f'النَّتَائِجُ الدِّرَاسِيَّةُ لِلْمُسْتَوَى {niveau}'))
    cv.setFont('DejaVuBold',12)
    cv.drawCentredString(W/2,y-83,f'Résultats Scolaires — Niveau {niveau}')
    cv.setFont('Amiri',12)
    cv.drawCentredString(W/2,y-97,ar('الدَّوْرَةُ الأُولَى')+'   •   Premier semestre')
    y -= HDR+5

    IR=24; cv.setFillColor(LGRY); cv.setStrokeColor(MAIN); cv.setLineWidth(1.2)
    cv.rect(MARGIN,y-IR,TW,IR,fill=1)
    cw4=TW/4
    nom_display = data.get('nom_ar','') or data.get('nom','')
    for i,txt in enumerate([ar(f'الاسم : {nom_display}'),ar(f'الفصل : {data.get("classe","")}'),ar(f'الأستاذ : {data.get("prof","")}'),ar('السنة : 2025-2026')]):
        if i>0: cv.setStrokeColor(MAIN); cv.setLineWidth(0.6); cv.line(MARGIN+i*cw4,y-IR,MARGIN+i*cw4,y)
        cv.setFont('Amiri',12); cv.setFillColor(MAIN)
        cv.drawCentredString(MARGIN+i*cw4+cw4/2,y-IR+7,txt)
    y -= IR+4

    y = th_row(y)
    y = sec_row(y,'اللُّغَةُ الْعَرَبِيَّةُ','LANGUE ARABE')
    nar=data.get('notes_arabe',[]); oar=data.get('obs_arabe',[])
    for i,(arm,frm) in enumerate(mats['arabe']):
        y=data_row(y,arm,frm,nar[i] if i<len(nar) else '',oar[i] if i<len(oar) else '',bg=LGRY if i%2==0 else WHITE)
    va=[float(n) for n in nar if n and str(n).replace('.','').isdigit()]
    aa=sum(va)/len(va) if va else None
    y=avg_row(y,'مُعَدَّلُ الْعَرَبِيَّةِ','Moy. Langue Arabe',aa)

    y=sec_row(y,'التَّرْبِيَةُ الْإِسْلَامِيَّةُ','ÉDUCATION ISLAMIQUE')
    ni=data.get('notes_islami',[]); oi=data.get('obs_islami',[])
    for i,(arm,frm) in enumerate(mats['islami']):
        y=data_row(y,arm,frm,ni[i] if i<len(ni) else '',oi[i] if i<len(oi) else '',bg=LGRY if i%2==0 else WHITE)
    vi=[float(n) for n in ni if n and str(n).replace('.','').isdigit()]
    ai=sum(vi)/len(vi) if vi else None
    y=avg_row(y,'مُعَدَّلُ الْإِسْلَامِيَّاتِ','Moy. Islamique',ai)

    y=sec_row(y,'الْمُرَاقَبَةُ الْمُسْتَمِرَّةُ','CONTRÔLE CONTINU')
    nc=''.join(c for c in str(data.get('note_continu','')) if c.isdigit() or c=='.')
    y=data_row(y,'الْمُرَاقَبَةُ الْمُسْتَمِرَّةُ','Contrôle Continu',nc,data.get('obs_continu',''),bg=LGRY)

    y=sec_row(y,'الْحُضُورُ وَالانْضِبَاطُ','PRÉSENCE & ASSIDUITÉ')
    np_=str(data.get('note_presence',''))
    y=data_row(y,'الْحُضُورُ وَالانْضِبَاطُ','Présence & Assiduité',np_,data.get('obs_presence',''),bg=LGRY)

    all_a=[]
    if aa is not None: all_a.append(aa)
    if ai is not None: all_a.append(ai)
    if nc and nc.replace('.','').isdigit(): all_a.append(float(nc))
    if np_ and np_.replace('.','').isdigit(): all_a.append(float(np_))
    ag=sum(all_a)/len(all_a) if all_a else None
    y=mg_row(y,ag)

    y-=6; ft=y; fb=MARGIN+16; fh=ft-fb
    cv.setStrokeColor(MAIN); cv.setLineWidth(1.5); cv.rect(MARGIN,fb,TW,fh)
    sd=min(fh-8,110); sw=sd+8
    xd1=MARGIN+sw; xd2=xd1+(TW-sw)/2
    cv.setStrokeColor(MAIN); cv.setLineWidth(0.8)
    cv.line(xd1,fb,xd1,ft); cv.line(xd2,fb,xd2,ft)
    stamp=ImageReader(STAMP_PATH)
    cv.drawImage(stamp,MARGIN+sw/2-sd/2,fb+fh/2-sd/2,width=sd,height=sd,mask='auto')
    P2=5
    cv.setFont('AmiriBold',11); cv.setFillColor(MAIN)
    cv.drawRightString(xd2-P2,ft-12,ar('مُلاحَظَاتُ الأُسْتَاذِ :'))
    op=data.get('obs_prof','')
    if op:
        lp=wrap_text(cv,op,'Amiri',10,(xd2-xd1)-2*P2)
        cv.setFont('Amiri',10); cv.setFillColor(BLK)
        for li,ln in enumerate(lp[:4]): cv.drawRightString(xd2-P2,ft-24-li*12,ar(ln))
    cv.setFont('DejaVu',8); cv.setFillColor(colors.HexColor('#444'))
    cv.drawString(xd1+P2,fb+10,'Signature : ________________')
    cv.setFont('AmiriBold',11); cv.setFillColor(MAIN)
    cv.drawRightString(X4-P2,ft-12,ar('مُلاحَظَاتُ الْوَالِدَيْنِ :'))
    op2=data.get('obs_parents','')
    if op2:
        lp2=wrap_text(cv,op2,'Amiri',10,(X4-xd2)-2*P2)
        cv.setFont('Amiri',10); cv.setFillColor(BLK)
        for li,ln in enumerate(lp2[:4]): cv.drawRightString(X4-P2,ft-24-li*12,ar(ln))
    cv.setFont('DejaVu',8); cv.setFillColor(colors.HexColor('#444'))
    cv.drawString(xd2+P2,fb+10,'Signature : ________________')
    cv.setFont('DejaVu',7); cv.setFillColor(colors.HexColor('#666'))
    cv.drawCentredString(W/2,MARGIN+4,'École Mosquée de Castelsarrasin  •  Association Musulmane de Castelsarrasin  •')


def calc_avg(notes):
    vals=[float(n) for n in notes if n and str(n).replace('.','').isdigit()]
    return round(sum(vals)/len(vals),2) if vals else None


def generate_excel(eleves_par_niveau):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    COLORS_NV = {
        1: {'header':'2e7d32','sub':'c8e6c9','title':'1b5e20'},
        2: {'header':'6a1b9a','sub':'e1bee7','title':'4a148c'},
        3: {'header':'1a2e5a','sub':'bbdefb','title':'0d1a38'},
    }

    for niveau, eleves in eleves_par_niveau.items():
        if not eleves: continue
        ws = wb.create_sheet(f"Niveau {niveau}")
        C = COLORS_NV[niveau]
        mats = MATIERES[niveau]
        mat_ar_labels = [m[1] for m in mats['arabe']]
        mat_isl_labels = [m[1] for m in mats['islami']]

        headers = ['Nom','Classe','Professeur'] + mat_ar_labels + ['Moy.Arabe'] + mat_isl_labels + ['Moy.Islamique','Ctrl.Continu','Présence','Moy.Générale','Obs.Prof']
        thin = Side(style='thin', color='CCCCCC')
        border = Border(left=thin,right=thin,top=thin,bottom=thin)

        # Title
        ws.merge_cells(f'A1:{get_column_letter(len(headers))}1')
        tc = ws['A1']
        tc.value = f"Bulletins Niveau {niveau} — École Mosquée de Castelsarrasin — 2025-2026"
        tc.font = Font(name='Calibri',bold=True,size=14,color='FFFFFF')
        tc.fill = PatternFill("solid",fgColor=C['title'])
        tc.alignment = Alignment(horizontal='center',vertical='center')
        ws.row_dimensions[1].height = 28

        # Headers
        for col,h in enumerate(headers,1):
            cell = ws.cell(row=2,column=col,value=h)
            cell.font = Font(name='Calibri',bold=True,size=10,color='FFFFFF')
            cell.fill = PatternFill("solid",fgColor=C['header'])
            cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
            cell.border = border
        ws.row_dimensions[2].height = 36

        # Data
        for row_idx,e in enumerate(eleves,3):
            nar=e.get('notes_arabe',[]); ni=e.get('notes_islami',[])
            nc=e.get('note_continu',''); np_=e.get('note_presence','')
            avg_ar=calc_avg(nar); avg_isl=calc_avg(ni)
            all_a=[x for x in [avg_ar,avg_isl,float(nc) if nc and nc.replace('.','').isdigit() else None,float(np_) if np_ and np_.replace('.','').isdigit() else None] if x is not None]
            avg_gen=round(sum(all_a)/len(all_a),2) if all_a else None

            row_data=[e.get('nom',''),e.get('classe',''),e.get('prof','')]
            row_data+=[nar[i] if i<len(nar) else '' for i in range(len(mat_ar_labels))]
            row_data+=[avg_ar or '']
            row_data+=[ni[i] if i<len(ni) else '' for i in range(len(mat_isl_labels))]
            row_data+=[avg_isl or '',nc,np_,avg_gen or '',e.get('obs_prof','')]

            bg = C['sub'] if row_idx%2==0 else 'FFFFFF'
            for col,val in enumerate(row_data,1):
                cell=ws.cell(row=row_idx,column=col,value=val)
                cell.fill=PatternFill("solid",fgColor=bg)
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                cell.border=border
                cell.font=Font(name='Calibri',size=10)

        # Column widths
        ws.column_dimensions['A'].width=20
        ws.column_dimensions['B'].width=14
        ws.column_dimensions['C'].width=18
        for i in range(4,len(headers)+1):
            ws.column_dimensions[get_column_letter(i)].width=14
        ws.column_dimensions[get_column_letter(len(headers))].width=30

        ws.freeze_panes='A3'

    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


# ── Routes ─────────────────────────────────────────────────

@app.route('/')
def index():
    return render_template('index.html', logo_uri=LOGO_URI)

@app.route('/api/generate_pdf', methods=['POST'])
def generate_pdf():
    data = request.json
    eleves = data.get('eleves', [])
    niveau = int(data.get('niveau', 3))
    merge = data.get('merge', False)
    inline = data.get('inline', False)   # True = ouvre dans le navigateur (impression)

    if not eleves:
        return jsonify({'error':'Aucun élève'}), 400

    pdfs = []
    for e in eleves:
        buf = io.BytesIO()
        c = pdfcanvas.Canvas(buf, pagesize=A4)
        draw_bulletin(c, e, niveau)
        c.save()
        buf.seek(0)
        pdfs.append(buf)

    as_attach = not inline
    if merge or len(pdfs)>1:
        writer=PdfWriter()
        for buf in pdfs:
            r=PdfReader(buf)
            for page in r.pages: writer.add_page(page)
        out=io.BytesIO(); writer.write(out); out.seek(0)
        fname=f"bulletins_niveau{niveau}.pdf"
        return send_file(out, mimetype='application/pdf', as_attachment=as_attach, download_name=fname)
    else:
        nom=eleves[0].get('nom','eleve').replace(' ','_')
        fname=f"bulletin_{nom}.pdf"
        return send_file(pdfs[0], mimetype='application/pdf', as_attachment=as_attach, download_name=fname)


@app.route('/api/generate_pdf_all', methods=['POST'])
def generate_pdf_all():
    """Génère un PDF fusionné pour tous les élèves de tous les niveaux."""
    data = request.json
    eleves_par_niveau = data.get('eleves', {})
    inline = data.get('inline', False)
    writer = PdfWriter()
    total = 0
    for niveau_str in ['1', '2', '3']:
        eleves = eleves_par_niveau.get(str(niveau_str), [])
        niveau = int(niveau_str)
        for e in eleves:
            buf = io.BytesIO()
            c = pdfcanvas.Canvas(buf, pagesize=A4)
            draw_bulletin(c, e, niveau)
            c.save()
            buf.seek(0)
            r = PdfReader(buf)
            for page in r.pages:
                writer.add_page(page)
            total += 1
    if total == 0:
        return jsonify({'error': 'Aucun élève dans aucun niveau'}), 400
    out = io.BytesIO()
    writer.write(out)
    out.seek(0)
    return send_file(out, mimetype='application/pdf',
                     as_attachment=not inline,
                     download_name='tous_bulletins.pdf')

@app.route('/api/generate_excel', methods=['POST'])
def generate_excel_route():
    data = request.json
    eleves_par_niveau = {}
    for niveau_str, eleves in data.items():
        try: eleves_par_niveau[int(niveau_str)] = eleves
        except: pass

    buf = generate_excel(eleves_par_niveau)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='bulletins_tous_niveaux.xlsx')

@app.route('/api/import_excel', methods=['POST'])
def import_excel():
    if 'file' not in request.files:
        return jsonify({'error':'Aucun fichier'}), 400
    f = request.files['file']
    niveau = int(request.form.get('niveau', 3))
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        mats = MATIERES[niveau]
        n_ar = len(mats['arabe']); n_isl = len(mats['islami'])
        eleves = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row[0]: continue
            nar=[str(row[3+i] or '') for i in range(n_ar)]
            ni=[str(row[4+n_ar+i] or '') for i in range(n_isl)]
            nc=str(row[5+n_ar+n_isl] or '')
            np_=str(row[6+n_ar+n_isl] or '')
            eleves.append({
                'nom':str(row[0] or ''),'classe':str(row[1] or ''),'prof':str(row[2] or ''),
                'notes_arabe':nar,'obs_arabe':['']*n_ar,
                'notes_islami':ni,'obs_islami':['']*n_isl,
                'note_continu':nc,'obs_continu':'',
                'note_presence':np_,'obs_presence':'',
                'obs_prof':str(row[7+n_ar+n_isl] or '') if len(row)>7+n_ar+n_isl else '',
                'obs_parents':''
            })
        return jsonify({'eleves':eleves,'count':len(eleves)})
    except Exception as ex:
        return jsonify({'error':str(ex)}), 500

@app.route('/api/import_json', methods=['POST'])
def import_json():
    if 'file' not in request.files:
        return jsonify({'error':'Aucun fichier'}), 400
    f = request.files['file']
    try:
        data = json.load(f)
        if isinstance(data, list): return jsonify({'eleves':data})
        if isinstance(data, dict) and 'eleves' in data: return jsonify({'eleves':data['eleves']})
        return jsonify({'error':'Format JSON non reconnu'}), 400
    except Exception as ex:
        return jsonify({'error':str(ex)}), 500

@app.route('/api/export_template_excel', methods=['GET'])
def export_template():
    niveau = int(request.args.get('niveau', 3))
    mats = MATIERES[niveau]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Niveau {niveau}"
    mat_ar=[m[1] for m in mats['arabe']]; mat_isl=[m[1] for m in mats['islami']]
    headers=['Nom','Classe','Professeur']+mat_ar+['Moy.Arabe']+mat_isl+['Moy.Islamique','Ctrl.Continu','Présence','Moy.Générale','Obs.Prof']
    thin=Side(style='thin',color='CCCCCC'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for col,h in enumerate(headers,1):
        cell=ws.cell(row=1,column=col,value=h)
        cell.font=Font(bold=True,size=10)
        cell.fill=PatternFill("solid",fgColor="1a2e5a")
        cell.font=Font(bold=True,size=10,color='FFFFFF')
        cell.alignment=Alignment(horizontal='center',wrap_text=True)
        cell.border=border
        ws.column_dimensions[get_column_letter(col)].width=16
    # 5 sample rows
    for r in range(2,7):
        for col in range(1,len(headers)+1):
            ws.cell(row=r,column=col,value='').border=border
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'template_niveau{niveau}.xlsx')

@app.route('/api/save_pdf', methods=['POST'])
def save_pdf():
    """Génère le PDF d'un élève et le sauvegarde dans exports/pdf/niv{n}/."""
    data = request.json
    eleve = data.get('eleve', {})
    niveau = int(data.get('niveau', 1))
    if not eleve.get('nom'):
        return jsonify({'error': 'Nom manquant'}), 400
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    draw_bulletin(c, eleve, niveau)
    c.save()
    buf.seek(0)
    nom_safe = safe_filename(eleve.get('nom', 'eleve'))
    filename = f"{nom_safe}.pdf"
    filepath = EXPORTS_DIR / f'niv{niveau}' / filename
    with open(filepath, 'wb') as fout:
        fout.write(buf.read())
    return jsonify({'success': True, 'filename': filename, 'niveau': niveau})

@app.route('/api/list_pdfs', methods=['GET'])
def list_pdfs():
    """Retourne la liste des PDFs sauvegardés par niveau."""
    result = {}
    for n in [1, 2, 3]:
        folder = EXPORTS_DIR / f'niv{n}'
        folder.mkdir(parents=True, exist_ok=True)
        files = sorted([f.name for f in folder.iterdir() if f.suffix == '.pdf'])
        result[str(n)] = files
    return jsonify(result)

@app.route('/api/get_pdf/<int:niveau>/<path:filename>', methods=['GET'])
def get_pdf(niveau, filename):
    """Télécharge un PDF sauvegardé."""
    filename = Path(filename).name  # sécurité anti path traversal
    filepath = EXPORTS_DIR / f'niv{niveau}' / filename
    if not filepath.exists():
        return jsonify({'error': 'Fichier non trouvé'}), 404
    return send_file(str(filepath), mimetype='application/pdf',
                     as_attachment=True, download_name=filename)

@app.route('/api/download_zip/<int:niveau>', methods=['GET'])
def download_zip(niveau):
    """Télécharge un ZIP de tous les PDFs (niveau 0 = tous les niveaux)."""
    if niveau == 0:
        folders = [(n, EXPORTS_DIR / f'niv{n}') for n in [1, 2, 3]]
        zip_name = 'bulletins_tous_niveaux.zip'
    else:
        folders = [(niveau, EXPORTS_DIR / f'niv{niveau}')]
        zip_name = f'bulletins_niv{niveau}.zip'
    buf = io.BytesIO()
    total = 0
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for n, folder in folders:
            if folder.exists():
                for pdf_file in sorted(folder.iterdir()):
                    if pdf_file.suffix == '.pdf':
                        zf.write(pdf_file, f'niv{n}/{pdf_file.name}')
                        total += 1
    if total == 0:
        return jsonify({'error': 'Aucun PDF sauvegardé'}), 404
    buf.seek(0)
    return send_file(buf, mimetype='application/zip',
                     as_attachment=True, download_name=zip_name)

if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=7860)
